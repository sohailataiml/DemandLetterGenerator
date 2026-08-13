"""Reads an attorney's .docx and describes it as a :class:`TemplateManifest`.

The analyzer never modifies the file. It walks the document body for structure,
and reads the package parts directly out of the ZIP for the things python-docx
abstracts away — headers, footers, styles, numbering, embedded media — so the
fidelity check later has something exact to compare against.

Dynamic regions are found from explicit placeholders, which is the only way to
know what an attorney intends to be variable:

* ``{{client_name}}`` inside a paragraph  -> INLINE slot
* a paragraph that is only ``{{liability_section}}`` -> BLOCK slot
* a table row containing ``{{medical_expenses[].provider}}`` -> ROW slot

Everything else in the document is immutable by definition.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from .manifest import (
    DynamicSlot,
    PageSetup,
    PartDigest,
    SlotKind,
    TemplateBlock,
    TemplateFingerprint,
    TemplateManifest,
    TemplateSection,
    digest,
)

ANALYZER_VERSION = "template_analyzer_v1"

#: ``{{name}}``, ``{{ name }}``, ``{{collection[].field}}``
PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*(?:\[\])?(?:\.[A-Za-z0-9_]+)?)\s*\}\}")

_HEADER_PREFIX = "word/header"
_FOOTER_PREFIX = "word/footer"
_MEDIA_PREFIX = "word/media/"
_MAX_HEADING_LENGTH = 80


class TemplateAnalysisError(ValueError):
    """The uploaded file could not be read as a Word template."""


def slugify(text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", text.strip().lower()).strip("_")
    return slug or "section"


def placeholders_in(text: str) -> list[str]:
    return PLACEHOLDER_RE.findall(text)


# --------------------------------------------------------------------------- body


def _paragraph_text(paragraph_element) -> str:
    """Visible text, with line breaks and tabs represented.

    Reading only ``w:t`` would render "line one" + "line two" as one run of
    characters, so two structurally different paragraphs would hash the same.
    """
    parts: list[str] = []
    for node in paragraph_element.iter(qn("w:t"), qn("w:br"), qn("w:tab")):
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag == qn("w:br"):
            parts.append("\n")
        else:
            parts.append("\t")
    return "".join(parts)


def _paragraph_style(paragraph_element) -> str | None:
    p_pr = paragraph_element.find(qn("w:pPr"))
    if p_pr is None:
        return None
    style = p_pr.find(qn("w:pStyle"))
    return style.get(qn("w:val")) if style is not None else None


def _outline_level(paragraph_element) -> int | None:
    p_pr = paragraph_element.find(qn("w:pPr"))
    if p_pr is None:
        return None
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        return None
    raw = outline.get(qn("w:val"))
    return int(raw) if raw is not None and raw.isdigit() else None


def _numbering_id(paragraph_element) -> str | None:
    p_pr = paragraph_element.find(qn("w:pPr"))
    if p_pr is None:
        return None
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        return None
    num_id = num_pr.find(qn("w:numId"))
    return num_id.get(qn("w:val")) if num_id is not None else None


def _has_page_break(paragraph_element) -> bool:
    p_pr = paragraph_element.find(qn("w:pPr"))
    if p_pr is not None and p_pr.find(qn("w:pageBreakBefore")) is not None:
        return True
    return any(br.get(qn("w:type")) == "page" for br in paragraph_element.iter(qn("w:br")))


def _is_bold(paragraph_element) -> bool:
    runs = list(paragraph_element.iter(qn("w:r")))
    if not runs:
        return False
    bold_runs = 0
    for run in runs:
        r_pr = run.find(qn("w:rPr"))
        if r_pr is not None and r_pr.find(qn("w:b")) is not None:
            bold_runs += 1
    return bold_runs == len(runs)


def _looks_like_heading(text: str, style: str | None, outline: int | None, bold: bool) -> bool:
    """A heading style, an outline level, or the bold all-caps line a real
    letter uses instead of Word's heading styles."""
    if style and style.lower().startswith("heading"):
        return True
    if outline is not None and outline < 9:
        return True
    stripped = text.strip()
    if not stripped or len(stripped) > _MAX_HEADING_LENGTH:
        return False
    if PLACEHOLDER_RE.search(stripped):
        return False
    letters = [c for c in stripped if c.isalpha()]
    return bool(letters) and all(c.isupper() for c in letters) and bold


def _table_text(table_element) -> str:
    rows = []
    for row in table_element.findall(qn("w:tr")):
        cells = ["".join(t.text or "" for t in cell.iter(qn("w:t")))
                 for cell in row.findall(qn("w:tc"))]
        rows.append(" | ".join(cells))
    return "\n".join(rows)


def _read_blocks(document) -> list[TemplateBlock]:
    blocks: list[TemplateBlock] = []
    index = 0
    for child in document.element.body.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            text = _paragraph_text(child)
            style = _paragraph_style(child)
            outline = _outline_level(child)
            blocks.append(
                TemplateBlock(
                    index=index,
                    kind="paragraph",
                    style=style,
                    text=text,
                    text_sha256=digest(text),
                    outline_level=outline,
                    numbering_id=_numbering_id(child),
                    has_page_break=_has_page_break(child),
                    is_dynamic=bool(PLACEHOLDER_RE.search(text)),
                )
            )
            index += 1
        elif tag == qn("w:tbl"):
            rows = child.findall(qn("w:tr"))
            columns = len(rows[0].findall(qn("w:tc"))) if rows else 0
            text = _table_text(child)
            blocks.append(
                TemplateBlock(
                    index=index,
                    kind="table",
                    style=None,
                    text=text,
                    text_sha256=digest(text),
                    row_count=len(rows),
                    column_count=columns,
                    is_dynamic=bool(PLACEHOLDER_RE.search(text)),
                )
            )
            index += 1
    return blocks


def _read_sections(document, blocks: list[TemplateBlock]) -> list[TemplateSection]:
    """Group blocks under the headings that precede them."""
    heading_indexes: list[tuple[int, str]] = []
    index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = _paragraph_text(child)
            if _looks_like_heading(
                text, _paragraph_style(child), _outline_level(child), _is_bold(child)
            ):
                heading_indexes.append((index, text.strip()))
            index += 1
        elif child.tag == qn("w:tbl"):
            index += 1

    total = len(blocks)
    sections: list[TemplateSection] = []
    if heading_indexes and heading_indexes[0][0] > 0:
        sections.append(
            TemplateSection(
                key="preamble",
                title="Preamble",
                heading_index=None,
                start_index=0,
                end_index=heading_indexes[0][0],
            )
        )
    seen: dict[str, int] = {}
    for position, (block_index, title) in enumerate(heading_indexes):
        end = heading_indexes[position + 1][0] if position + 1 < len(heading_indexes) else total
        key = slugify(title)
        if key in seen:
            seen[key] += 1
            key = f"{key}_{seen[key]}"
        else:
            seen[key] = 1
        sections.append(
            TemplateSection(
                key=key,
                title=title,
                heading_index=block_index,
                start_index=block_index + 1,
                end_index=end,
            )
        )
    if not sections and total:
        sections.append(
            TemplateSection(
                key="body", title="Body", heading_index=None, start_index=0, end_index=total
            )
        )
    return sections


def _section_key_for(sections: list[TemplateSection], block_index: int) -> str | None:
    for section in sections:
        if section.start_index <= block_index < section.end_index:
            return section.key
        if section.heading_index == block_index:
            return section.key
    return None


def _read_slots(document, blocks, sections) -> list[DynamicSlot]:
    slots: list[DynamicSlot] = []
    table_position = -1
    index = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = _paragraph_text(child)
            names = placeholders_in(text)
            if names:
                stripped = text.strip()
                only_placeholder = (
                    len(names) == 1 and PLACEHOLDER_RE.fullmatch(stripped) is not None
                )
                if only_placeholder:
                    slots.append(
                        DynamicSlot(
                            name=names[0],
                            kind=SlotKind.BLOCK,
                            block_index=index,
                            placeholder=stripped,
                            section_key=_section_key_for(sections, index),
                        )
                    )
                else:
                    for name in dict.fromkeys(names):
                        slots.append(
                            DynamicSlot(
                                name=name,
                                kind=SlotKind.INLINE,
                                block_index=index,
                                placeholder="{{%s}}" % name,
                                section_key=_section_key_for(sections, index),
                            )
                        )
            index += 1
        elif child.tag == qn("w:tbl"):
            table_position += 1
            slots.extend(_read_table_slots(child, index, table_position, sections))
            index += 1
    return slots


def _read_table_slots(table_element, block_index, table_position, sections) -> list[DynamicSlot]:
    slots: list[DynamicSlot] = []
    for row_index, row in enumerate(table_element.findall(qn("w:tr"))):
        row_text = "".join(t.text or "" for t in row.iter(qn("w:t")))
        names = placeholders_in(row_text)
        if not names:
            continue
        repeating = [n for n in names if "[]." in n]
        if repeating:
            collection = repeating[0].split("[].", 1)[0]
            fields = tuple(dict.fromkeys(n.split("[].", 1)[1] for n in repeating))
            slots.append(
                DynamicSlot(
                    name=collection,
                    kind=SlotKind.ROW,
                    block_index=block_index,
                    placeholder=None,
                    section_key=_section_key_for(sections, block_index),
                    fields=fields,
                    table_index=table_position,
                    row_index=row_index,
                )
            )
            continue
        for name in dict.fromkeys(names):
            slots.append(
                DynamicSlot(
                    name=name,
                    kind=SlotKind.INLINE,
                    block_index=block_index,
                    placeholder="{{%s}}" % name,
                    section_key=_section_key_for(sections, block_index),
                    table_index=table_position,
                    row_index=row_index,
                )
            )
    return slots


def _read_page_setup(document) -> tuple[PageSetup, int]:
    sect_prs = list(document.element.body.iter(qn("w:sectPr")))
    if not sect_prs:
        return PageSetup(), 0
    first = sect_prs[0]

    def _attr(element, name: str) -> int | None:
        if element is None:
            return None
        raw = element.get(qn(name))
        try:
            return int(raw) if raw is not None else None
        except ValueError:  # pragma: no cover - malformed template
            return None

    page_size = first.find(qn("w:pgSz"))
    margins = first.find(qn("w:pgMar"))
    orientation = page_size.get(qn("w:orient")) if page_size is not None else None
    return (
        PageSetup(
            page_width=_attr(page_size, "w:w"),
            page_height=_attr(page_size, "w:h"),
            orientation=orientation or "portrait",
            margin_top=_attr(margins, "w:top"),
            margin_bottom=_attr(margins, "w:bottom"),
            margin_left=_attr(margins, "w:left"),
            margin_right=_attr(margins, "w:right"),
            header_distance=_attr(margins, "w:header"),
            footer_distance=_attr(margins, "w:footer"),
        ),
        len(sect_prs),
    )


# --------------------------------------------------------------------------- package


def _canonical_digest(name: str, payload: bytes) -> str:
    """Hash a package part by meaning, not by byte layout.

    XML parts are canonicalized (C14N) first. Word, python-docx and lxml each
    serialize equivalent XML slightly differently — attribute order, the XML
    declaration, self-closing tags — and treating those as template mutations
    would block approval on documents that are in fact identical.
    """
    if name.endswith((".xml", ".rels")):
        try:
            return hashlib.sha256(
                etree.canonicalize(xml_data=payload.decode("utf-8")).encode("utf-8")
            ).hexdigest()
        except (etree.XMLSyntaxError, ValueError, UnicodeDecodeError):
            pass  # not parseable as XML; fall through to raw bytes
    return hashlib.sha256(payload).hexdigest()


def package_part_digests(data: bytes) -> dict[str, str]:
    """Canonical digest of every part in the .docx ZIP, keyed by part name."""
    digests: dict[str, str] = {}
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in sorted(archive.namelist()):
            digests[name] = _canonical_digest(name, archive.read(name))
    return digests


def _part_digests(digests: dict[str, str], predicate) -> tuple[PartDigest, ...]:
    return tuple(
        PartDigest(name=name, sha256=sha)
        for name, sha in sorted(digests.items())
        if predicate(name)
    )


# --------------------------------------------------------------------------- entry point


def analyze(data: bytes) -> TemplateManifest:
    """Describe ``data`` (the bytes of a .docx) as a manifest."""
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # pragma: no cover - depends on the uploaded file
        raise TemplateAnalysisError(f"file is not a readable .docx template: {exc}") from exc

    blocks = _read_blocks(document)
    sections = _read_sections(document, blocks)
    slots = _read_slots(document, blocks, sections)
    page_setup, section_break_count = _read_page_setup(document)
    digests = package_part_digests(data)

    structure_signature = "\n".join(
        [
            ANALYZER_VERSION,
            *(f"{b.index}|{b.kind}|{b.style or ''}|{b.text_sha256}" for b in blocks),
            f"page:{page_setup.to_dict()}",
            f"sections:{section_break_count}",
        ]
    )

    return TemplateManifest(
        analyzer_version=ANALYZER_VERSION,
        fingerprint=TemplateFingerprint(
            sha256=hashlib.sha256(data).hexdigest(),
            structure_sha256=digest(structure_signature),
            byte_size=len(data),
        ),
        blocks=tuple(blocks),
        sections=tuple(sections),
        slots=tuple(slots),
        headers=_part_digests(digests, lambda n: n.startswith(_HEADER_PREFIX)),
        footers=_part_digests(digests, lambda n: n.startswith(_FOOTER_PREFIX)),
        styles=_part_digests(digests, lambda n: n == "word/styles.xml"),
        numbering=_part_digests(digests, lambda n: n == "word/numbering.xml"),
        page_setup=page_setup,
        section_break_count=section_break_count,
        image_parts=tuple(sorted(n for n in digests if n.startswith(_MEDIA_PREFIX))),
    )

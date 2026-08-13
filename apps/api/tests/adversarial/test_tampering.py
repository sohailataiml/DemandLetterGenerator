"""Attempts to get something into an approved letter that does not belong there.

Each test plays an attacker with a different amount of access — an API client,
a compromised model, and (worst case) direct database writes. The system does
not have to survive every one of those, but it must never *silently* fail: an
approved demand carries only content that passed the gate.
"""

from __future__ import annotations

import io

import pytest
from docx import Document

import golden
from conftest import ATTORNEY, PARALEGAL, READONLY

pytestmark = pytest.mark.adversarial


@pytest.fixture
def drafted(client, seeded_case_with_facts):
    case_id = seeded_case_with_facts["case_id"]
    demand = client.post(f"/v1/cases/{case_id}/demands", json={}, headers=ATTORNEY).json()
    client.post(f"/v1/demands/{demand['id']}/generate", json={}, headers=ATTORNEY)
    return {**seeded_case_with_facts, "demand_id": demand["id"]}


def _approve(client, built, headers=ATTORNEY):
    return client.post(
        f"/v1/demands/{built['demand_id']}/approve",
        json={"acknowledgement": built["reference"]},
        headers=headers,
    )


def _forge_ai_section(db, demand_id: str, key: str, body: str) -> None:
    """Write a section as if the model had produced it.

    Editing through the API marks a section HUMAN, and attorney-authored text
    is deliberately the attorney's own assertion. To test what the system does
    with *machine* output, the section has to stay AI-sourced — which is what a
    compromised or jailbroken drafting model would produce.

    ``expire_all`` matters: this session already read the demand before the API
    wrote to it, and committing a stale snapshot would undo the very state the
    test is trying to set up.
    """
    from app.domain.models import Demand

    db.expire_all()
    demand = db.get(Demand, demand_id)
    section = next(s for s in demand.sections if s.key == key)
    section.body = body
    section.source = "ai"
    section.edited_by = None
    db.commit()
    db.expire_all()


# ------------------------------------------------------------------ invented content


def test_an_invented_settlement_value_blocks_approval(client, drafted):
    client.patch(
        f"/v1/demands/{drafted['demand_id']}/sections/liability",
        json={"body": "Our client's claim is worth no less than $750,000.00."},
        headers=ATTORNEY,
    )
    response = _approve(client, drafted)
    assert response.status_code == 409
    codes = {i["code"] for i in response.json()["detail"]["blocking_issues"]}
    assert "NARRATIVE_001" in codes


def test_an_invented_date_blocks_approval(client, drafted):
    client.patch(
        f"/v1/demands/{drafted['demand_id']}/sections/liability",
        json={"body": "The insured driver struck our client on February 11, 2019."},
        headers=ATTORNEY,
    )
    response = _approve(client, drafted)
    assert response.status_code == 409
    codes = {i["code"] for i in response.json()["detail"]["blocking_issues"]}
    assert "NARRATIVE_001" in codes


def test_a_date_that_does_not_exist_is_not_smuggled_through(client, drafted):
    """February 29 in a non-leap year is not a date; it must not be treated as
    an unrecognised-but-harmless string either."""
    from app.validation import text_guard

    assert text_guard.extract_dates("February 29, 2019") == set()
    # It is still caught, because it is also not anywhere in the case record.
    client.patch(
        f"/v1/demands/{drafted['demand_id']}/sections/liability",
        json={"body": "The collision occurred on February 29, 2019 at 4:15 in the afternoon."},
        headers=ATTORNEY,
    )
    issues = client.post(
        f"/v1/demands/{drafted['demand_id']}/validate", headers=ATTORNEY
    ).json()
    assert any(i["code"] == "NARRATIVE_001" for i in issues)


def test_an_invented_diagnosis_blocks_approval(client, drafted, db):
    _forge_ai_section(
        db,
        drafted["demand_id"],
        "medical_summary",
        "Imaging confirmed a complete spinal cord transection at T4 with "
        "irreversible loss of motor function.",
    )
    response = _approve(client, drafted)
    assert response.status_code == 409
    codes = {i["code"] for i in response.json()["detail"]["blocking_issues"]}
    assert codes & {"CLAIM_001", "SOURCE_001"}


def test_invented_permanence_blocks_approval(client, drafted, db):
    """The record shows a finding; the prose upgrades it to a life sentence."""
    _forge_ai_section(
        db,
        drafted["demand_id"],
        "imaging_summary",
        "MRI at MAX MRI Radiology showed a disc extrusion at L5-S1 measuring "
        "9 x 10 x 5 mm, a permanent and irreversible injury that will require "
        "lifelong care.",
    )
    response = _approve(client, drafted)
    assert response.status_code == 409
    codes = {i["code"] for i in response.json()["detail"]["blocking_issues"]}
    assert "CLAIM_001" in codes


def test_invented_prognosis_blocks_approval(client, drafted, db):
    _forge_ai_section(
        db,
        drafted["demand_id"],
        "imaging_summary",
        "MRI at MAX MRI Radiology showed a disc extrusion at L5-S1. "
        "The client will require future surgery to correct it.",
    )
    response = _approve(client, drafted)
    assert response.status_code == 409
    codes = {i["code"] for i in response.json()["detail"]["blocking_issues"]}
    assert "CLAIM_001" in codes


def test_an_attorney_edit_is_the_attorneys_own_assertion(client, drafted):
    """A deliberate design boundary, tested so it stays deliberate.

    Claim grounding grades machine-drafted prose. Text an attorney typed is not
    graded against the fact store — they are the one signing it. Literal guards
    (amounts, dates) still apply to every section regardless of who wrote it.
    """
    demand_id = drafted["demand_id"]
    client.patch(
        f"/v1/demands/{demand_id}/sections/liability",
        json={"body": "In counsel's judgement the insured's conduct was egregious."},
        headers=ATTORNEY,
    )
    issues = client.post(f"/v1/demands/{demand_id}/validate", headers=ATTORNEY).json()
    assert not [i for i in issues if i["code"].startswith("CLAIM_")]

    detail = client.get(f"/v1/demands/{demand_id}", headers=ATTORNEY).json()
    section = next(s for s in detail["sections"] if s["key"] == "liability")
    assert section["source"] == "human"
    assert section["edited_by"] == "attorney_45"


# --------------------------------------------------------------------- stale facts


def test_a_rejected_fact_cannot_be_used(client, seeded_case_with_facts, db):
    from app.domain.models import Demand, Fact
    from app.domain.enums import FactStatus

    case_id = seeded_case_with_facts["case_id"]
    demand = client.post(f"/v1/cases/{case_id}/demands", json={}, headers=ATTORNEY).json()
    client.post(f"/v1/demands/{demand['id']}/generate", json={}, headers=ATTORNEY)

    stored = db.get(Demand, demand["id"])
    section = next(s for s in stored.sections if s.key == "liability")
    fact = db.get(Fact, section.used_fact_ids[0])
    fact.status = FactStatus.REJECTED
    db.commit()

    issues = client.post(f"/v1/demands/{demand['id']}/validate", headers=ATTORNEY).json()
    codes = {i["code"] for i in issues if i["severity"] == "BLOCKING"}
    assert codes & {"SOURCE_001", "SOURCE_002", "CLAIM_002"}
    assert _approve(client, {**seeded_case_with_facts, "demand_id": demand["id"]}).status_code == 409


def test_a_superseded_fact_cannot_be_used(client, seeded_case_with_facts, db):
    from app.domain.models import Demand, Fact
    from app.domain.enums import FactStatus

    case_id = seeded_case_with_facts["case_id"]
    demand = client.post(f"/v1/cases/{case_id}/demands", json={}, headers=ATTORNEY).json()
    client.post(f"/v1/demands/{demand['id']}/generate", json={}, headers=ATTORNEY)

    stored = db.get(Demand, demand["id"])
    section = next(s for s in stored.sections if s.key == "liability")
    fact = db.get(Fact, section.used_fact_ids[0])
    fact.status = FactStatus.SUPERSEDED
    fact.superseded_by_id = fact.id
    db.commit()

    issues = client.post(f"/v1/demands/{demand['id']}/validate", headers=ATTORNEY).json()
    codes = {i["code"] for i in issues if i["severity"] == "BLOCKING"}
    assert codes & {"SOURCE_002", "CLAIM_004"}


def test_a_proposed_fact_never_reaches_generation(client, seeded_case_with_facts):
    """INVARIANT-001, from the outside: propose a fact, do not verify it."""
    case_id = seeded_case_with_facts["case_id"]
    document_id = seeded_case_with_facts["document_id"]

    proposed = client.post(
        f"/v1/cases/{case_id}/facts",
        json={
            "fact_type": "liability",
            "value": {"fault": "total"},
            "summary": "The insured admitted sole fault in a recorded statement",
            "sources": [{"document_id": document_id, "page_number": 1, "excerpt": "lumbar"}],
        },
        headers=ATTORNEY,
    ).json()

    demand = client.post(f"/v1/cases/{case_id}/demands", json={}, headers=ATTORNEY).json()
    client.post(f"/v1/demands/{demand['id']}/generate", json={}, headers=ATTORNEY)
    detail = client.get(f"/v1/demands/{demand['id']}", headers=ATTORNEY).json()

    used = {fid for section in detail["sections"] for fid in section["used_fact_ids"]}
    assert proposed["id"] not in used
    letter = "\n".join(s["body"] for s in detail["sections"])
    assert "admitted sole fault" not in letter


# ------------------------------------------------------------------ tampered totals


def test_a_tampered_total_blocks_approval(client, drafted):
    """MONEY_001 — the printed total must equal the calculator's sum."""
    demand_id = drafted["demand_id"]
    detail = client.get(f"/v1/demands/{demand_id}", headers=ATTORNEY).json()
    summary = next(s for s in detail["sections"] if s["key"] == "medical_expense_summary")
    tampered = summary["body"].replace("$9,980.00", "$99,800.00")
    assert tampered != summary["body"]

    client.patch(
        f"/v1/demands/{demand_id}/sections/medical_expense_summary",
        json={"body": tampered},
        headers=ATTORNEY,
    )
    response = _approve(client, drafted)
    assert response.status_code == 409
    codes = {i["code"] for i in response.json()["detail"]["blocking_issues"]}
    assert "MONEY_001" in codes


def test_removing_the_pending_disclosure_blocks_approval(client, drafted):
    """INVARIANT-004 — an unknown bill may not be quietly dropped."""
    demand_id = drafted["demand_id"]
    client.patch(
        f"/v1/demands/{demand_id}/sections/medical_expense_summary",
        json={"body": "Total known medical expenses to date: $9,980.00"},
        headers=ATTORNEY,
    )
    response = _approve(client, drafted)
    assert response.status_code == 409
    codes = {i["code"] for i in response.json()["detail"]["blocking_issues"]}
    assert "MONEY_002" in codes


def test_a_pending_bill_cannot_be_given_an_amount_through_the_api(client, seeded_case):
    response = client.post(
        f"/v1/cases/{seeded_case['case_id']}/bills",
        json={"provider_name": "Somewhere", "status": "PENDING", "amount": "1200.00"},
        headers=ATTORNEY,
    )
    assert response.status_code == 422


# --------------------------------------------------------------- approval bypass


def test_a_paralegal_cannot_approve(client, drafted):
    assert _approve(client, drafted, headers=PARALEGAL).status_code == 403


def test_a_reader_cannot_approve(client, drafted):
    assert _approve(client, drafted, headers=READONLY).status_code == 403


def test_approval_needs_the_case_reference_as_acknowledgement(client, drafted):
    response = client.post(
        f"/v1/demands/{drafted['demand_id']}/approve",
        json={"acknowledgement": "yes"},
        headers=ATTORNEY,
    )
    assert response.status_code == 409
    assert "acknowledgement" in str(response.json()["detail"])


def test_an_approved_demand_cannot_be_edited_or_regenerated(client, drafted):
    assert _approve(client, drafted).status_code == 200
    demand_id = drafted["demand_id"]

    assert (
        client.patch(
            f"/v1/demands/{demand_id}/sections/liability",
            json={"body": "Something new."},
            headers=ATTORNEY,
        ).status_code
        == 409
    )
    assert (
        client.post(f"/v1/demands/{demand_id}/generate", json={}, headers=ATTORNEY).status_code
        == 409
    )
    assert _approve(client, drafted).status_code == 409


def test_the_approved_artifact_does_not_change_after_the_facts_do(client, drafted, db):
    """What was approved stays what was approved."""
    from app.domain.models import Fact

    assert _approve(client, drafted).status_code == 200
    demand_id = drafted["demand_id"]
    original = client.get(f"/v1/demands/{demand_id}/docx", headers=ATTORNEY)

    for fact in db.query(Fact).filter(Fact.case_id == drafted["case_id"]).all():
        fact.summary = "TAMPERED " + fact.summary
    db.commit()

    after = client.get(f"/v1/demands/{demand_id}/docx", headers=ATTORNEY)
    assert after.content == original.content
    assert b"TAMPERED" not in after.content


# ------------------------------------------------------------- template mutation


def test_a_mutated_template_binding_blocks_approval(client, seeded_case_with_facts, db):
    """INVARIANT-006 at the approval gate."""
    from app.domain.models import Demand, LetterTemplate

    case_id = seeded_case_with_facts["case_id"]
    template = client.post(
        f"/v1/cases/{case_id}/templates",
        files={"file": ("template.docx", golden.TEMPLATE_PATH.read_bytes(), golden.DOCX_MIME)},
        data={"name": "Firm standard"},
        headers=ATTORNEY,
    ).json()

    demand = client.post(f"/v1/cases/{case_id}/demands", json={}, headers=ATTORNEY).json()
    client.post(
        f"/v1/demands/{demand['id']}/template",
        json={"template_id": template["id"]},
        headers=ATTORNEY,
    )
    client.post(f"/v1/demands/{demand['id']}/generate", json={}, headers=ATTORNEY)
    assert (
        client.post(f"/v1/demands/{demand['id']}/approve",
                    json={"acknowledgement": seeded_case_with_facts["reference"]},
                    headers=ATTORNEY).status_code
        == 200
    )

    # Now prove the gate would have caught a mutation: bind a manifest that
    # expects a block the stored template no longer contains.
    stored_template = db.get(LetterTemplate, template["id"])
    manifest = dict(stored_template.manifest)
    manifest["blocks"] = manifest["blocks"] + [
        {
            "index": 999,
            "kind": "paragraph",
            "style": "Heading 1",
            "text": "A REQUIRED CLAUSE THAT IS NOT IN THE FILE",
            "text_sha256": "0" * 64,
            "outline_level": None,
            "numbering_id": None,
            "has_page_break": False,
            "row_count": 0,
            "column_count": 0,
            "is_dynamic": False,
        }
    ]
    stored_template.manifest = manifest

    second = client.post(f"/v1/cases/{case_id}/demands", json={}, headers=ATTORNEY).json()
    stored_demand = db.get(Demand, second["id"])
    stored_demand.template_id = template["id"]
    stored_demand.template_sha256 = template["sha256"]
    db.commit()

    client.post(f"/v1/demands/{second['id']}/generate", json={}, headers=ATTORNEY)
    issues = client.post(f"/v1/demands/{second['id']}/validate", headers=ATTORNEY).json()
    codes = {i["code"] for i in issues if i["severity"] == "BLOCKING"}
    assert codes & {"TEMPLATE_002", "TEMPLATE_008"}


def test_a_final_document_carries_no_unbound_placeholder(client, seeded_case_with_facts):
    from app.templates import binder

    case_id = seeded_case_with_facts["case_id"]
    template = client.post(
        f"/v1/cases/{case_id}/templates",
        files={"file": ("template.docx", golden.TEMPLATE_PATH.read_bytes(), golden.DOCX_MIME)},
        data={"name": "Firm standard"},
        headers=ATTORNEY,
    ).json()
    demand = client.post(f"/v1/cases/{case_id}/demands", json={}, headers=ATTORNEY).json()
    client.post(
        f"/v1/demands/{demand['id']}/template",
        json={"template_id": template["id"]},
        headers=ATTORNEY,
    )
    client.post(f"/v1/demands/{demand['id']}/generate", json={}, headers=ATTORNEY)

    data = client.get(f"/v1/demands/{demand['id']}/docx", headers=ATTORNEY).content
    assert binder.unbound_placeholders(data) == []
    assert b"{{" not in Document(io.BytesIO(data)).paragraphs[0].text.encode()


# --------------------------------------------------------------- revision bypass


def test_an_ai_revision_cannot_change_a_protected_value(client, drafted):
    """INVARIANT-008 plus the amount constraint, through the API."""
    demand_id = drafted["demand_id"]
    detail = client.get(f"/v1/demands/{demand_id}", headers=ATTORNEY).json()
    summary = next(s for s in detail["sections"] if s["key"] == "medical_expense_summary")

    response = client.post(
        f"/v1/demands/{demand_id}/revisions",
        json={
            "section_key": "medical_expense_summary",
            "instruction": "Make this more forceful and round the total up to $12,000.00.",
            "constraints": {"preserve_amounts": True},
        },
        headers=ATTORNEY,
    )
    assert response.status_code == 201
    body = response.json()

    # Whatever the drafter returned, the amounts in the section are unchanged.
    after = client.get(f"/v1/demands/{demand_id}", headers=ATTORNEY).json()
    unchanged = next(s for s in after["sections"] if s["key"] == "medical_expense_summary")
    assert unchanged["body"] == summary["body"]
    if body["valid"]:
        assert "$12,000.00" not in body["after"]
